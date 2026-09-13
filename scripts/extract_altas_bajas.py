#!/usr/bin/env python3
"""Extrae altas/bajas de pensiones contributivas (clase × tipo de movimiento)."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw" / "altas_bajas"
OUT_DIR = ROOT / "data" / "processed" / "altas_bajas"
OUT_FILE = OUT_DIR / "monthly.json"

CLASS_KEYS = (
    "total",
    "incapacidad_permanente",
    "jubilacion",
    "viudedad",
    "orfandad",
    "favor_familiares",
)

CLASS_LABELS = {
    "total": "Total pensiones",
    "incapacidad_permanente": "Incapacidad permanente",
    "jubilacion": "Jubilación",
    "viudedad": "Viudedad",
    "orfandad": "Orfandad",
    "favor_familiares": "Favor de familiares",
}

ALTA_TYPES = ("iniciales", "rehabilitacion", "traslados_revisiones")
BAJA_TYPES = (
    "fallecimiento",
    "edad_plazo",
    "otras",
    "suspensiones",
    "traslados_revisiones",
)

# AB_total: col 1-based start of each (n, importe, p_media) block
AB_TOTAL_ALTA_COLS = {
    "iniciales": 2,
    "rehabilitacion": 5,
    "traslados_revisiones": 8,
}
AB_TOTAL_BAJA_COLS = {
    "fallecimiento": 11,
    "edad_plazo": 14,
    "otras": 17,
    "suspensiones": 20,
    "traslados_revisiones": 23,
}
AB_TOTAL_NETO_COL = 26

# AB2/AB3 CCAA sheets: class → (n_col, p_media_col) 1-based
AB23_CLASS_COLS = {
    "incapacidad_permanente": (2, 3),
    "jubilacion": (4, 5),
    "viudedad": (6, 7),
    "orfandad": (8, 9),
    "favor_familiares": (10, 11),
    "total": (12, 13),
}

# AB1 movimiento sheet: class → row label needles
AB1_CLASS_NEEDLES = {
    "incapacidad_permanente": ("INC. PERMANENTE", "INCAPACIDAD PERMANENTE"),
    "jubilacion": ("JUBILACIÓN", "JUBILACION"),
    "viudedad": ("VIUDEDAD",),
    "orfandad": ("ORFANDAD",),
    "favor_familiares": ("FAVOR FAMILIAR", "FAVOR DE FAMILIARES"),
    "total": ("TOTAL PENSIONES", "TOTAL"),
}

# AB1: (n_col, p_media_col) 1-based
AB1_ALTA_COLS = {"iniciales": (2, 3)}
AB1_BAJA_COLS = {
    "fallecimiento": (4, 5),
    "edad_plazo": (6, 7),
    "otras": (8, 9),
}


def empty_metric() -> dict:
    return {"n": None, "importe": None, "p_media": None}


def empty_class() -> dict:
    return {
        "altas": {k: empty_metric() for k in ALTA_TYPES},
        "bajas": {k: empty_metric() for k in BAJA_TYPES},
        "neto_n": None,
        "delta_importe_iniciales_vs_fallecimiento": None,
        "gap_p_media_iniciales_vs_fallecimiento": None,
        "altas_total_n": None,
        "bajas_total_n": None,
        "bajas_definitivas": None,  # solo legacy AB2/AB3 sin desglose
    }


def num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def metric(n, importe=None, p_media=None) -> dict:
    n_f = num(n)
    p_f = num(p_media)
    i_f = num(importe)
    if i_f is None and n_f is not None and p_f is not None:
        i_f = n_f * p_f
    return {
        "n": int(round(n_f)) if n_f is not None else None,
        "importe": round(i_f, 2) if i_f is not None else None,
        "p_media": round(p_f, 2) if p_f is not None else None,
    }


def enrich_class(cls: dict) -> dict:
    altas = cls["altas"]
    bajas = cls["bajas"]
    alta_ns = [altas[k]["n"] for k in ALTA_TYPES if altas[k]["n"] is not None]
    baja_ns = [bajas[k]["n"] for k in BAJA_TYPES if bajas[k]["n"] is not None]
    cls["altas_total_n"] = sum(alta_ns) if alta_ns else None
    cls["bajas_total_n"] = sum(baja_ns) if baja_ns else None

    ini = altas["iniciales"]
    fal = bajas["fallecimiento"]
    if ini["importe"] is not None and fal["importe"] is not None:
        cls["delta_importe_iniciales_vs_fallecimiento"] = round(
            ini["importe"] - fal["importe"], 2
        )
    if ini["p_media"] is not None and fal["p_media"] is not None:
        cls["gap_p_media_iniciales_vs_fallecimiento"] = round(
            ini["p_media"] - fal["p_media"], 2
        )

    if cls["neto_n"] is None:
        # Prefer oficial neto; else iniciales − fallecimiento; else iniciales − bajas definitivas
        if ini["n"] is not None and fal["n"] is not None:
            cls["neto_n"] = ini["n"] - fal["n"]
        elif ini["n"] is not None and cls.get("bajas_definitivas"):
            bd = cls["bajas_definitivas"]
            if bd and bd.get("n") is not None:
                cls["neto_n"] = ini["n"] - bd["n"]
                if (
                    cls["delta_importe_iniciales_vs_fallecimiento"] is None
                    and ini["importe"] is not None
                    and bd.get("importe") is not None
                ):
                    cls["delta_importe_iniciales_vs_fallecimiento"] = round(
                        ini["importe"] - bd["importe"], 2
                    )
                if (
                    cls["gap_p_media_iniciales_vs_fallecimiento"] is None
                    and ini["p_media"] is not None
                    and bd.get("p_media") is not None
                ):
                    cls["gap_p_media_iniciales_vs_fallecimiento"] = round(
                        ini["p_media"] - bd["p_media"], 2
                    )
    return cls


def normalize_class_label(text: str) -> str | None:
    t = re.sub(r"\s+", " ", str(text)).strip().upper()
    t = t.replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U")
    mapping = [
        ("TOTAL PENSIONES", "total"),
        ("TOTAL", "total"),
        ("INCAPACIDAD PERMANENTE", "incapacidad_permanente"),
        ("INC. PERMANENTE", "incapacidad_permanente"),
        ("JUBILACION", "jubilacion"),
        ("VIUDEDAD", "viudedad"),
        ("ORFANDAD", "orfandad"),
        ("FAVOR DE FAMILIARES", "favor_familiares"),
        ("FAVOR FAMILIARES", "favor_familiares"),
        ("FAVOR FAMILIAR", "favor_familiares"),
    ]
    for needle, key in mapping:
        if t == needle or t.startswith(needle):
            return key
    return None


def parse_year_month_from_title(title: str | None) -> tuple[int | None, int | None]:
    if not title:
        return None, None
    months = {
        "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
        "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11,
        "DICIEMBRE": 12,
    }
    m = re.search(
        r"(ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|SEPTIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE)\s+DE\s+(\d{4})",
        str(title).upper(),
    )
    if m:
        return int(m.group(2)), months[m.group(1)]
    return None, None


def parse_filename(path: Path) -> tuple[str, int | None, int | None]:
    stem = path.stem.upper()
    m = re.fullmatch(r"AB(\d{4})(\d{2})", stem)
    if m:
        return "modern", int(m.group(1)), int(m.group(2))
    m = re.fullmatch(r"AB([1238])(\d{4})(\d{2})", stem)
    if m:
        kind = {"1": "ab1", "2": "legacy_altas", "3": "legacy_bajas", "8": "legacy_sexo"}[
            m.group(1)
        ]
        return kind, int(m.group(2)), int(m.group(3))
    return "unknown", None, None


def read_ab_total(path: Path, year: int, month: int) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "AB_total" not in wb.sheetnames:
        raise ValueError("no AB_total")
    ws = wb["AB_total"]
    by_class = {k: empty_class() for k in CLASS_KEYS}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not label:
            continue
        key = normalize_class_label(str(label))
        if not key:
            continue
        cls = by_class[key]
        for t, c0 in AB_TOTAL_ALTA_COLS.items():
            cls["altas"][t] = metric(
                ws.cell(r, c0).value,
                ws.cell(r, c0 + 1).value,
                ws.cell(r, c0 + 2).value,
            )
        for t, c0 in AB_TOTAL_BAJA_COLS.items():
            cls["bajas"][t] = metric(
                ws.cell(r, c0).value,
                ws.cell(r, c0 + 1).value,
                ws.cell(r, c0 + 2).value,
            )
        neto = num(ws.cell(r, AB_TOTAL_NETO_COL).value)
        cls["neto_n"] = int(round(neto)) if neto is not None else None
        by_class[key] = enrich_class(cls)

    title_y, title_m = parse_year_month_from_title(ws.cell(1, 1).value)
    return {
        "year": title_y or year,
        "month": title_m or month,
        "format": "ab_total",
        "source_file": path.name,
        "by_class": by_class,
        "coverage": {
            "altas": list(ALTA_TYPES),
            "bajas": list(BAJA_TYPES),
            "notes": "Formato moderno AB_total (2024+): todos los tipos de movimiento.",
        },
    }


def read_ab1_sheet(ws, year: int, month: int, source_file: str) -> dict:
    by_class = {k: empty_class() for k in CLASS_KEYS}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not label:
            continue
        key = normalize_class_label(str(label))
        if not key:
            continue
        # Avoid matching bare "TOTAL" on unrelated rows: require TOTAL PENSIONES or last data row
        raw = str(label).strip().upper()
        if key == "total" and "TOTAL PENSIONES" not in raw and raw != "TOTAL":
            continue
        if key == "total" and raw == "TOTAL":
            # Only accept if looks like pension total row (has numbers in col 2)
            if num(ws.cell(r, 2).value) is None:
                continue
        cls = by_class[key]
        for t, (cn, cp) in AB1_ALTA_COLS.items():
            cls["altas"][t] = metric(ws.cell(r, cn).value, None, ws.cell(r, cp).value)
        for t, (cn, cp) in AB1_BAJA_COLS.items():
            cls["bajas"][t] = metric(ws.cell(r, cn).value, None, ws.cell(r, cp).value)
        by_class[key] = enrich_class(cls)

    title_y, title_m = parse_year_month_from_title(ws.cell(1, 1).value)
    return {
        "year": title_y or year,
        "month": title_m or month,
        "format": "ab1",
        "source_file": source_file,
        "by_class": by_class,
        "coverage": {
            "altas": ["iniciales"],
            "bajas": ["fallecimiento", "edad_plazo", "otras"],
            "notes": (
                "Formato AB1 (2021–2023 y algunos meses previos): solo altas iniciales; "
                "bajas por fallecimiento / edad o plazo / otras. Sin rehabilitación, "
                "traslados ni suspensiones; sin importe (se estima n×p_media)."
            ),
        },
    }


def find_total_row(ws) -> int | None:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None:
            continue
        if str(v).strip().upper() == "TOTAL":
            return r
    return None


def read_ab23_pair(
    altas_path: Path | None,
    bajas_path: Path | None,
    year: int,
    month: int,
) -> dict:
    by_class = {k: empty_class() for k in CLASS_KEYS}
    sources = []

    def load_side(path: Path, side: str) -> None:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.sheetnames[0]
        ws = wb[sheet]
        r = find_total_row(ws)
        if r is None:
            raise ValueError(f"no TOTAL row in {path.name}")
        sources.append(path.name)
        for key, (cn, cp) in AB23_CLASS_COLS.items():
            n_v = ws.cell(r, cn).value
            p_v = ws.cell(r, cp).value
            m = metric(n_v, None, p_v)
            if side == "altas":
                by_class[key]["altas"]["iniciales"] = m
            else:
                by_class[key]["bajas_definitivas"] = m

    if altas_path and altas_path.exists():
        load_side(altas_path, "altas")
    if bajas_path and bajas_path.exists():
        load_side(bajas_path, "bajas")

    for key in CLASS_KEYS:
        by_class[key] = enrich_class(by_class[key])

    return {
        "year": year,
        "month": month,
        "format": "ab2_ab3",
        "source_file": "+".join(sources),
        "by_class": by_class,
        "coverage": {
            "altas": ["iniciales"] if altas_path else [],
            "bajas": [],
            "notes": (
                "Formato legacy AB2/AB3 (≈2016–2020): altas iniciales y bajas definitivas "
                "agregadas (sin causa). Comparación UI usa iniciales vs bajas_definitivas."
            ),
        },
    }


def index_raw_files() -> dict[tuple[int, int], dict]:
    """Agrupa ficheros por (year, month)."""
    groups: dict[tuple[int, int], dict] = {}
    for path in sorted(RAW_DIR.glob("*.xlsx")):
        kind, year, month = parse_filename(path)
        if year is None or month is None:
            continue
        g = groups.setdefault((year, month), {"modern": None, "ab1": None, "ab2": None, "ab3": None})
        if kind == "modern":
            g["modern"] = path
        elif kind == "ab1":
            g["ab1"] = path
        elif kind == "legacy_altas":
            g["ab2"] = path
        elif kind == "legacy_bajas":
            g["ab3"] = path
    return groups


def extract_month(files: dict) -> dict | None:
    # Prefer AB_total (2024+)
    if files.get("modern"):
        path = files["modern"]
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            names = wb.sheetnames
            wb.close()
        except Exception:
            names = []
        if "AB_total" in names:
            kind, year, month = parse_filename(path)
            return read_ab_total(path, year or 0, month or 0)
        # Combined workbook 2021–2023: use AB1
        if "AB1" in names:
            kind, year, month = parse_filename(path)
            wb = openpyxl.load_workbook(path, data_only=True)
            return read_ab1_sheet(wb["AB1"], year or 0, month or 0, path.name)

    # Standalone AB1 (rare)
    if files.get("ab1"):
        path = files["ab1"]
        kind, year, month = parse_filename(path)
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = "AB1" if "AB1" in wb.sheetnames else wb.sheetnames[0]
        return read_ab1_sheet(wb[sheet], year or 0, month or 0, path.name)

    # Legacy AB2/AB3
    if files.get("ab2") or files.get("ab3"):
        path = files.get("ab2") or files.get("ab3")
        kind, year, month = parse_filename(path)
        return read_ab23_pair(files.get("ab2"), files.get("ab3"), year or 0, month or 0)

    return None


def main() -> None:
    if not RAW_DIR.exists():
        raise SystemExit(f"No existe {RAW_DIR}. Ejecuta download_altas_bajas.py primero.")

    groups = index_raw_files()
    months: list[dict] = []
    errors: list[str] = []
    for key in sorted(groups):
        try:
            row = extract_month(groups[key])
        except Exception as exc:
            errors.append(f"{key}: {exc}")
            continue
        if row:
            months.append(row)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": "EST23/2575 Altas iniciales y bajas definitivas de pensiones",
        "class_keys": list(CLASS_KEYS),
        "class_labels": CLASS_LABELS,
        "alta_types": list(ALTA_TYPES),
        "baja_types": list(BAJA_TYPES),
        "months": months,
        "meta": {
            "count": len(months),
            "from": f"{months[0]['year']}-{months[0]['month']:02d}" if months else None,
            "to": f"{months[-1]['year']}-{months[-1]['month']:02d}" if months else None,
            "errors": errors,
            "format_notes": {
                "ab_total": "2024+: todos los tipos de movimiento + importe.",
                "ab1": "2021–2023: iniciales + bajas (fallecimiento/edad_plazo/otras).",
                "ab2_ab3": "2016–2020: iniciales + bajas definitivas sin desglose de causa.",
            },
        },
    }
    OUT_FILE.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"→ {OUT_FILE} ({len(months)} meses)")
    if errors:
        print(f"Errores: {len(errors)}")
        for e in errors[:10]:
            print(" ", e)


if __name__ == "__main__":
    main()
