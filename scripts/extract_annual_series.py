#!/usr/bin/env python3
"""Extrae serie anual agregada desde los XLSX ResSISTEMA (cierre de ejercicio)."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

RAW_DIR = Path(__file__).resolve().parents[1] / "data" / "raw" / "res_sistema"
OUT_FILE = Path(__file__).resolve().parents[1] / "data" / "processed" / "annual_series.json"

MONTH_ORDER = {
    "Enero": 1,
    "Febrero": 2,
    "Marzo": 3,
    "Abril": 4,
    "Mayo": 5,
    "Junio": 6,
    "Julio": 7,
    "Agosto": 8,
    "Septiembre": 9,
    "Octubre": 10,
    "Noviembre": 11,
    "Diciembre": 12,
    "DiciembreProvisional": 13,
    "DiciembreDefinitivo": 14,
}


def parse_filename(name: str) -> tuple[str | None, int | None]:
    stem = name.replace("ResSISTEMA_", "").replace(".xlsx", "")
    for pattern in (
        r"^(Diciembre(?:Provisional|Definitivo)?|[A-Za-zÁÉÍÓÚñÑ]+)_(\d{4})$",
        r"^(Diciembre(?:Provisional|Definitivo)?|[A-Za-zÁÉÍÓÚñÑ]+)(\d{4})$",
        r"^(DiciembreDefinitivo)_(\d{4})D$",
    ):
        match = re.match(pattern, stem)
        if match:
            return match.group(1), int(match.group(2))
    return None, None


def find_col(rows: list, *parts: str) -> int:
    for i, row in enumerate(rows[:20]):
        for j, cell in enumerate(row):
            if cell is None:
                continue
            text = str(cell).upper().replace("\n", " ")
            if all(part in text for part in parts):
                return j
            if "RECAUDACI" in text:
                nxt = rows[i + 1][j] if i + 1 < len(rows) and j < len(rows[i + 1]) else None
                if nxt and "NETA" in str(nxt).upper():
                    return j
    return 3


def find_row(rows: list, *needles: str):
    for row in rows:
        if row and row[0] and any(needle in str(row[0]).upper() for needle in needles):
            return row
    return None


def val(row, col: int) -> float | None:
    if row is None or col >= len(row):
        return None
    value = row[col]
    return float(value) if isinstance(value, (int, float)) else None


def pick_recaudacion(row) -> float | None:
    if not row:
        return None
    money = [float(v) for v in row if isinstance(v, (int, float)) and float(v) > 500]
    if len(money) >= 3:
        return money[2]
    if len(money) >= 2:
        return money[1]
    return money[0] if money else None


def extract_file(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pi = list(workbook["ResumenPI"].iter_rows(values_only=True))
    pic4 = list(workbook["ResumenPICap4"].iter_rows(values_only=True))
    pg = list(workbook["ResumenPG"].iter_rows(values_only=True))

    pag_col = find_col(pg, "PAGOS", "REALIZAD") or 7

    cot_row = find_row(pi, "1. COTIZACIONES SOCIALES")
    tr_row = find_row(pi, "4. TRANSFERENCIAS CORRIENTES")
    ing_row = find_row(pi, "TOTAL INGRESOS")
    est_row = find_row(pic4, "DEL ESTADO Y ORG.AUTÓNOMOS")
    gast_row = find_row(pg, "TOTAL GASTOS")

    gast = val(gast_row, pag_col)
    if gast is None and gast_row:
        gast = max(float(v) for v in gast_row if isinstance(v, (int, float)))

    return {
        "cotizaciones": pick_recaudacion(cot_row),
        "transferencias": pick_recaudacion(tr_row),
        "transferencias_estado": pick_recaudacion(est_row),
        "ingresos_totales": pick_recaudacion(ing_row),
        "gastos": gast,
    }


def pick_year_end_files() -> dict[int, tuple[int, str, Path]]:
    by_year: dict[int, tuple[int, str, Path]] = {}
    for path in RAW_DIR.glob("ResSISTEMA*.xlsx"):
        month, year = parse_filename(path.name)
        if not year:
            continue
        priority = MONTH_ORDER.get(month, 0)
        if year not in by_year or priority > by_year[year][0]:
            by_year[year] = (priority, month, path)
    return by_year


def main() -> None:
    records = []
    for year, (_, month, path) in sorted(pick_year_end_files().items()):
        if year >= 2026:
            continue
        metrics = extract_file(path)
        if not metrics["cotizaciones"] or not metrics["gastos"]:
            continue
        metrics = {key: round(value, 1) if value is not None else None for key, value in metrics.items()}
        metrics["year"] = year
        metrics["source_month"] = month
        metrics["deficit_cotizaciones"] = round(metrics["gastos"] - metrics["cotizaciones"], 1)
        if metrics["transferencias_estado"] and metrics["ingresos_totales"]:
            metrics["pct_estado_en_ingresos"] = round(
                metrics["transferencias_estado"] / metrics["ingresos_totales"] * 100, 1
            )
        else:
            metrics["pct_estado_en_ingresos"] = None
        records.append(metrics)

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUT_FILE.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {len(records)} years to {OUT_FILE}")


if __name__ == "__main__":
    main()
