#!/usr/bin/env python3
"""Extrae desglose anual de ingresos, gastos y transferencias desde ResSISTEMA."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw" / "res_sistema"
OUT_FILE = ROOT / "data" / "processed" / "annual_breakdown.json"
MONTHLY_DIR = ROOT / "data" / "processed" / "monthly"
MONTHLY_MANIFEST = MONTHLY_DIR / "manifest.json"

MONTH_ORDER = {
    "Enero": 1, "Febrero": 2, "Marzo": 3, "Abril": 4, "Mayo": 5, "Junio": 6,
    "Julio": 7, "Agosto": 8, "Septiembre": 9, "Octubre": 10, "Noviembre": 11,
    "Diciembre": 12, "DiciembreProvisional": 13, "DiciembreDefinitivo": 14,
}

# (needle en col A, clave json)
INGRESOS_PI = [
    ("1. COTIZACIONES SOCIALES", "cotizaciones"),
    ("3. TASAS Y OTROS INGRESOS", "tasas_otros"),
    ("4. TRANSFERENCIAS CORRIENTES", "transferencias_corrientes"),
    ("5. INGRESOS PATRIMONIALES", "ingresos_patrimoniales"),
    ("TOTAL INGRESOS", "ingresos_totales"),
]

INGRESOS_CAP4 = [
    ("RÉGIMEN GENERAL", "cot_regimen_general"),
    ("RÉGIMEN ESPECIAL TRABAJADORES AUTÓNOMOS", "cot_autonomos"),
    ("ACCIDENTES TRABAJO ENFERMEDADES PROFESIONALES", "cot_at_ep"),
    ("DESEMPLEADOS Y BENEF", "cot_desempleo"),
    ("MECANISMO DE EQUIDAD INTERGENERACIONAL", "cot_mei"),
    ("TOTAL COTIZACIONES SOCIALES", "cotizaciones_cap4"),
    ("DEL ESTADO Y ORG.AUTÓNOMOS", "transf_estado"),
    ("DE LA SEGURIDAD SOCIAL", "transf_ss_internas"),
    ("DE COMUNIDADES AUTÓNOMAS", "transf_ccaa"),
    ("DE EMPRESAS PRIVADAS Y OTROS", "transf_privados"),
    ("TOTAL TRANSFERENCIAS CORRIENTES", "transferencias_corrientes_cap4"),
]

GASTOS_PG = [
    ("1. GASTOS DE PERSONAL", "gasto_personal"),
    ("2. GASTOS CORRIENTES BIENES Y SERVICIOS", "gasto_bienes_servicios"),
    ("4. TRANSFERENCIAS CORRIENTES", "gasto_transferencias"),
    ("6. INVERSIONES REALES", "gasto_inversiones"),
    ("TOTAL GASTOS", "gastos_totales"),
]

GASTOS_CAP4 = [
    ("SUMA DE PENSIONES", "pensiones"),
    ("JUBILACIÓN", "pension_jubilacion"),
    ("VIUDEDAD", "pension_viudedad"),
    ("INCAPACIDAD PERMANENTE", "pension_incap_permanente"),
    ("SUMA DE SUBSIDIOS Y OTRAS PRESTACIONES", "subsidios_prestaciones"),
    ("INCAPACIDAD TEMPORAL", "prest_it"),
    ("SUBSIDIO TEMPORAL POR CONTINGENCIAS COMUNES", "prest_desempleo"),
    ("A LA SEGURIDAD SOCIAL", "gasto_transf_ss"),
    ("A COMUNIDADES AUTÓNOMAS", "gasto_transf_ccaa"),
    ("AL ESTADO Y ORG.AUTÓNOMOS", "gasto_transf_estado"),
]

GASTOS_NOCONT = [
    ("SUMA DE PENSIONES NO CONTR", "pensiones_no_contributivas"),
    ("COMPL. A MÍNIMOS PENSIONES CONTRIBUTIVAS", "complementos_minimos"),
    ("INGRESO MÍNIMO VITAL", "imv"),
    ("SUMA SUBSIDIOS Y OTRAS PRESTACIONES", "prest_no_cont_otras"),
    ("TOTAL", "gastos_no_contributivos"),
]


def parse_filename(name: str) -> tuple[str | None, int | None]:
    stem = name.replace("ResSISTEMA_", "").replace(".xlsx", "")
    for pattern in (
        r"^(Diciembre(?:Provisional|Definitivo)?|[A-Za-zÁÉÍÓÚñÑ]+)_(\d{4})$",
        r"^(Diciembre(?:Provisional|Definitivo)?|[A-Za-zÁÉÍÓÚñÑ]+)(\d{4})$",
        r"^(DiciembreDefinitivo)_(\d{4})D$",
    ):
        match = re.match(pattern, stem)
        if match:
            return match.group(1), int(match.group(2))
    return None, None


def find_col(rows: list, *parts: str) -> int:
    for i, row in enumerate(rows[:20]):
        for j, cell in enumerate(row):
            if cell is None:
                continue
            text = str(cell).upper().replace("\n", " ")
            if all(part in text for part in parts):
                return j
            if "RECAUDACI" in text:
                nxt = rows[i + 1][j] if i + 1 < len(rows) and j < len(rows[i + 1]) else None
                if nxt and "NETA" in str(nxt).upper():
                    return j
    return 6


def find_pagos_col(rows: list) -> int:
    for i, row in enumerate(rows[:12]):
        for j, cell in enumerate(row):
            if cell and "PAGOS REALIZADOS" in str(cell).upper().replace("\n", " "):
                return j
    return 7


def row_val(row, col: int) -> float | None:
    if not row or col >= len(row):
        return None
    v = row[col]
    return float(v) if isinstance(v, (int, float)) else None


def pick_recaudacion(row) -> float | None:
    """Recaudación neta: 3.er importe >500 M€ de la fila (layout ResSISTEMA)."""
    if not row:
        return None
    money = [float(v) for v in row if isinstance(v, (int, float)) and float(v) > 500]
    if len(money) >= 3:
        return money[2]
    if len(money) >= 2:
        return money[1]
    return money[0] if money else None


def extract_map(rows: list, col: int, mapping: list[tuple[str, str]]) -> dict[str, float | None]:
    out: dict[str, float | None] = {}
    for needle, key in mapping:
        row = find_row(rows, needle)
        out[key] = row_val(row, col)
    return out


def extract_map_recaudacion(rows: list, mapping: list[tuple[str, str]]) -> dict[str, float | None]:
    """Para ResumenPICap4: la columna de recaudación varía según el año."""
    out: dict[str, float | None] = {}
    for needle, key in mapping:
        row = find_row(rows, needle)
        out[key] = pick_recaudacion(row)
    return out


def find_row(rows: list, needle: str):
    needle = needle.upper()
    for row in rows:
        if row and row[0] and needle in str(row[0]).upper():
            return row
    return None


def extract_file(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pi = list(wb["ResumenPI"].iter_rows(values_only=True))
    pic4 = list(wb["ResumenPICap4"].iter_rows(values_only=True))
    pg = list(wb["ResumenPG"].iter_rows(values_only=True))
    pgcap4 = list(wb["ResumenPGCap4"].iter_rows(values_only=True))
    pgnoc = list(wb["ResumenPGCap4NoCont"].iter_rows(values_only=True))

    rec_col = find_col(pi, "RECAUDACI", "NETA")
    pag_col = find_pagos_col(pg)

    data: dict[str, float | None] = {}
    data.update(extract_map(pi, rec_col, INGRESOS_PI))
    data.update(extract_map_recaudacion(pic4, INGRESOS_CAP4))
    data.update(extract_map(pg, pag_col, GASTOS_PG))
    data.update(extract_map(pgcap4, pag_col, GASTOS_CAP4))
    data.update(extract_map(pgnoc, pag_col, GASTOS_NOCONT))

    # Otros regímenes de cotización = total - desglose principal
    cot_total = data.get("cotizaciones") or data.get("cotizaciones_cap4")
    if cot_total:
        known = sum(
            v or 0
            for k, v in data.items()
            if k.startswith("cot_") and k not in ("cotizaciones", "cotizaciones_cap4") and v
        )
        if known < cot_total:
            data["cot_otros_regimenes"] = round(cot_total - known, 1)

    return data


def pick_year_end_files() -> dict[int, tuple[int, str, Path]]:
    by_year: dict[int, tuple[int, str, Path]] = {}
    for path in RAW_DIR.glob("ResSISTEMA*.xlsx"):
        month, year = parse_filename(path.name)
        if not year:
            continue
        priority = MONTH_ORDER.get(month, 0)
        if year not in by_year or priority > by_year[year][0]:
            by_year[year] = (priority, month, path)
    return by_year


def round_record(d: dict) -> dict:
    return {k: round(v, 1) if isinstance(v, float) else v for k, v in d.items()}


def enrich_record(rec: dict) -> dict:
    """Métricas derivadas comunes a anual y mensual."""
    cot = rec.get("cotizaciones") or 0
    gast = rec.get("gastos_totales") or 0
    est = rec.get("transf_estado") or 0
    ing = rec.get("ingresos_totales") or 0
    rec["deficit_cotizaciones"] = round(gast - cot, 1) if gast and cot else None
    rec["cobertura_cotizaciones_pct"] = round(cot / gast * 100, 1) if gast else None
    rec["pct_estado_en_ingresos"] = round(est / ing * 100, 1) if ing and est else None
    rec["transf_estado_neta_aparente"] = round(est - (rec.get("gasto_transf_estado") or 0), 1)
    return rec


def monthly_json_path(year: int, month: str) -> Path:
    return MONTHLY_DIR / str(year) / f"{month}.json"


def load_monthly(year: int, month: str) -> dict | None:
    path = monthly_json_path(year, month)
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def latest_monthly(year: int, *, calendar_only: bool = True) -> tuple[dict | None, str | None]:
    """Último mes disponible en JSON para un año."""
    year_dir = MONTHLY_DIR / str(year)
    if not year_dir.is_dir():
        return None, None
    best_prio, best_month, best_rec = 0, None, None
    for path in year_dir.glob("*.json"):
        month = path.stem
        if calendar_only and month not in MONTH_ORDER:
            continue
        prio = MONTH_ORDER.get(month, 0)
        if prio > best_prio:
            best_prio, best_month = prio, month
            best_rec = json.loads(path.read_text(encoding="utf-8"))
    return best_rec, best_month


def list_monthly_year(year: int, *, calendar_only: bool = True) -> list[dict]:
    year_dir = MONTHLY_DIR / str(year)
    if not year_dir.is_dir():
        return []
    rows: list[dict] = []
    for path in year_dir.glob("*.json"):
        month = path.stem
        if calendar_only and month not in MONTH_ORDER:
            continue
        rec = json.loads(path.read_text(encoding="utf-8"))
        rec.setdefault("month", month)
        rec.setdefault("year", year)
        rows.append(rec)
    rows.sort(key=lambda r: MONTH_ORDER.get(r["month"], 0))
    return rows


def export_monthly_snapshots() -> int:
    """Extrae todos los xlsx locales a JSON mensual + manifest."""
    from datetime import datetime, timezone

    MONTHLY_DIR.mkdir(parents=True, exist_ok=True)
    entries: list[dict] = []
    latest_by_year: dict[str, str] = {}
    count = 0

    for path in sorted(RAW_DIR.glob("ResSISTEMA*.xlsx")):
        month, year = parse_filename(path.name)
        if not month or not year:
            continue
        raw = extract_file(path)
        if not raw.get("cotizaciones") or not raw.get("gastos_totales"):
            continue
        rec = enrich_record(round_record(raw))
        rec["year"] = year
        rec["month"] = month
        rec["source_file"] = path.name

        out_path = monthly_json_path(year, month)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(json.dumps(rec, indent=2, ensure_ascii=False), encoding="utf-8")
        count += 1

        entries.append({"year": year, "month": month, "file": f"{year}/{month}.json"})
        if month in MONTH_ORDER:
            cur = latest_by_year.get(str(year))
            if not cur or MONTH_ORDER[month] > MONTH_ORDER.get(cur, 0):
                latest_by_year[str(year)] = month

    manifest = {
        "source": "ResSISTEMA xlsx → JSON (extract_breakdown.py)",
        "updated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "count": count,
        "latest_by_year": latest_by_year,
        "entries": entries,
    }
    MONTHLY_MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {count} monthly snapshots to {MONTHLY_DIR}")
    return count


def main() -> None:
    export_monthly_snapshots()

    records = []
    for year, (_, month, path) in sorted(pick_year_end_files().items()):
        if year < 2014 or year >= 2026:
            continue
        raw = extract_file(path)
        if not raw.get("cotizaciones") or not raw.get("gastos_totales"):
            continue
        rec = enrich_record(round_record(raw))
        rec["year"] = year
        rec["source_month"] = month
        records.append(rec)

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUT_FILE.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {len(records)} years to {OUT_FILE}")


if __name__ == "__main__":
    main()
